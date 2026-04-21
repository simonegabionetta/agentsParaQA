#!/usr/bin/env python3
"""Converte arquivos casos_de_teste_*.md para .docx usando python-docx."""

import re
import glob
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


PRIORITY_COLORS = {
    "alta":  RGBColor(0xC0, 0x00, 0x00),
    "média": RGBColor(0xC5, 0x5A, 0x11),
    "media": RGBColor(0xC5, 0x5A, 0x11),
    "baixa": RGBColor(0x37, 0x86, 0x10),
}


def shade_cell(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def add_paragraph(doc, text: str, style="Normal", bold=False, size=None, color=None, align=None):
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if align:
        p.alignment = align
    return p


def parse_table(lines: list[str]):
    """Extrai linhas de uma tabela markdown."""
    rows = []
    for line in lines:
        if re.match(r"^\s*\|[-| :]+\|\s*$", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def render_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    headers = rows[0]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"

    # Cabeçalho
    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        shade_cell(cell, "D9D9D9")
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)

    # Dados
    for row_data in rows[1:]:
        row = table.add_row()
        for i, val in enumerate(row_data):
            if i < len(row.cells):
                cell = row.cells[i]
                cell.text = val
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
                        # Colorir coluna de prioridade
                        key = val.lower().strip()
                        if key in PRIORITY_COLORS:
                            run.font.color.rgb = PRIORITY_COLORS[key]
                            run.bold = True

    doc.add_paragraph()


def md_to_docx(md_path: str, docx_path: str):
    text = Path(md_path).read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()

    # Margens
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2)
        section.right_margin = Inches(1.2)

    i = 0
    table_lines = []
    in_table = False

    while i < len(lines):
        line = lines[i]

        # Linha de tabela
        if line.strip().startswith("|"):
            if not in_table:
                in_table = True
                table_lines = []
            table_lines.append(line)
            i += 1
            continue

        # Fim de bloco de tabela
        if in_table:
            render_table(doc, parse_table(table_lines))
            in_table = False
            table_lines = []

        # Headings
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        # Separador
        elif re.match(r"^---+$", line.strip()):
            doc.add_paragraph("─" * 60)
        # Bullets
        elif line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, line[2:])
        # Linha em branco
        elif line.strip() == "":
            doc.add_paragraph()
        # Parágrafo normal
        else:
            p = doc.add_paragraph()
            _add_inline(p, line)

        i += 1

    if in_table:
        render_table(doc, parse_table(table_lines))

    doc.save(docx_path)
    print(f"  ✅ {docx_path}")


def _add_inline(p, text: str):
    """Adiciona texto com suporte a **bold** inline."""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)


def main():
    files = sorted(glob.glob("casos_de_teste_*.md"))
    if not files:
        print("❌ Nenhum arquivo casos_de_teste_*.md encontrado.")
        sys.exit(1)

    print(f"\n📄 Convertendo {len(files)} arquivo(s)...\n")
    for md_path in files:
        docx_path = md_path.replace(".md", ".docx")
        md_to_docx(md_path, docx_path)

    print(f"\n✅ Arquivos .docx gerados em: {Path('.').resolve()}\n")


if __name__ == "__main__":
    main()
